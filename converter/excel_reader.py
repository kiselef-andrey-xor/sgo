# -*- coding: utf-8 -*-
"""Разбор таблицы расписания в формате Excel.

Ожидаемая структура листа (как в «расписание по предметам.xlsx»):

    A1 «Дни» | B1 «Уроки» | C1.. «1а», «1б», ... — названия классов
    A2 «Пн»  | B2 0        | C2 «Разговоры о важном (105)» ...

Ячейка = «Предмет (кабинет)»; для деления на группы —
«Иностранный язык (английский) (305,307)» (один предмет, два кабинета)
или «Информатика, Труд (технология) (320,110)» (два предмета, два кабинета).
"""
from __future__ import annotations

import datetime as _dt
import re
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

from .text_utils import norm

DAY_WORDS = {
    "пн": "понедельник", "пон": "понедельник", "пнд": "понедельник", "понедельник": "понедельник",
    "вт": "вторник", "втр": "вторник", "вторник": "вторник",
    "ср": "среда", "среда": "среда", "сред": "среда",
    "чт": "четверг", "четв": "четверг", "четверг": "четверг",
    "пт": "пятница", "пятн": "пятница", "пятница": "пятница",
    "сб": "суббота", "суб": "суббота", "суббота": "суббота",
    "вс": "воскресенье", "воскр": "воскресенье", "воскресенье": "воскресенье",
}
WEEKDAY_NAMES = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]

_CELL_RE = re.compile(r"^(?P<body>.*)[\(\[]\s*(?P<rooms>[^\)\]]*)\s*[\)\]]\s*$", re.DOTALL)
# похоже ли содержимое последних скобок на кабинеты: «305», «спорт.зал», «IT куб 1», «БАС»…
_ROOM_WORDS = ("зал", "спорт", "куб", "бас", "цпд", "каб", "стадион", "басс", "полиг",
               "мастерск", "плац", "фок", "офис", "блок")


def _looks_like_rooms(text: str) -> bool:
    parts = [p.strip() for p in text.split(",") if p.strip()]
    if not parts:
        return False
    for part in parts:
        low = part.lower()
        if re.search(r"\d", part):
            continue
        if any(w in low for w in _ROOM_WORDS):
            continue
        if part.isupper() and len(part) <= 4:
            continue
        return False
    return True
_TIME_RE = re.compile(r"^\s*(\d{1,2})[:.ч](\d{2})")


@dataclass
class Pair:
    """Одна пара «предмет + кабинет» внутри ячейки."""
    subject: str
    room: str
    cell: str
    row: int
    col: int
    raw: str


@dataclass
class Entry:
    """Содержимое одной ячейки расписания."""
    day_raw: str
    day: str                 # нормализованное название дня
    period_raw: str          # что написано в колонке «Уроки»
    period: Optional[int]    # номер урока (0 — «нулевой» урок)
    time_text: Optional[str] # если в колонке время «08:30»
    excel_class: str         # название класса из шапки
    col_index: int
    row_index: int
    pairs: List[Pair] = field(default_factory=list)
    raw: str = ""


@dataclass
class ExcelSchedule:
    path: str
    sheet: str
    header_row: int
    day_col: int
    lesson_col: int
    class_cols: List[Tuple[int, str]] = field(default_factory=list)
    entries: List[Entry] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


def split_top_level(text: str, sep: str = ",") -> List[str]:
    """Разбивает строку по разделителю вне скобок.

    «Иностранный язык (английский), Труд (технология)» ->
    ['Иностранный язык (английский)', 'Труд (технология)']
    """
    parts: List[str] = []
    buf = ""
    depth = 0
    for ch in text:
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth = max(0, depth - 1)
        if ch == sep and depth == 0:
            parts.append(buf)
            buf = ""
        else:
            buf += ch
    parts.append(buf)
    return [p.strip(" \t\r\n;·•-–—") for p in parts if p.strip(" \t\r\n;·•-–—")]


def parse_cell(text: str, cell: str, row: int, col: int) -> List[Pair]:
    """Разбирает текст ячейки на пары «предмет — кабинет»."""
    text = str(text).replace("\r", " ").strip()
    if not text:
        return []
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    pairs: List[Pair] = []
    for line in lines:
        m = _CELL_RE.match(line)
        if m and _looks_like_rooms(m.group("rooms")):
            body, rooms_txt = m.group("body").strip(), m.group("rooms").strip()
            subjects = split_top_level(body)
            rooms = split_top_level(rooms_txt)
        else:
            subjects, rooms = split_top_level(line), []
        if not subjects:
            continue
        if len(subjects) == 1 and len(rooms) > 1:
            # деление на группы: один предмет, несколько кабинетов
            pairs += [Pair(subjects[0], r, cell, row, col, line) for r in rooms]
        elif len(rooms) == 1 and len(subjects) > 1:
            # несколько предметов в одном кабинете (совмещённый урок)
            pairs += [Pair(s, rooms[0] if i == 0 else "", cell, row, col, line)
                      for i, s in enumerate(subjects)]
        else:
            n = max(len(subjects), len(rooms))
            subjects += [""] * (n - len(subjects))
            rooms += [""] * (n - len(rooms))
            pairs += [Pair(s, r, cell, row, col, line) for s, r in zip(subjects, rooms)]
    return pairs


def parse_period(value, lesson_times=None) -> Tuple[Optional[int], Optional[str], str]:
    """Номер урока из колонки «Уроки».

    Поддерживает числа (0, 1, 2…), текст «1 урок», время «08:30» и
    диапазон «08:30-09:10» (время сверяется со звонками эталона).
    """
    if value is None:
        return None, None, ""
    if isinstance(value, (_dt.time, _dt.datetime)):
        text = value.strftime("%H:%M")
        return _period_from_time(text, lesson_times), text, str(value)
    text = str(value).strip()
    if not text:
        return None, None, ""
    m = _TIME_RE.match(text)
    if m:
        hhmm = f"{int(m.group(1)):02d}:{m.group(2)}"
        return _period_from_time(hhmm, lesson_times), hhmm, text
    digits = re.sub(r"\D", "", text.split("-")[0].split("–")[0])
    if digits == "":
        return None, None, text
    return int(digits), None, text


def _period_from_time(hhmm: str, lesson_times) -> Optional[int]:
    if not lesson_times:
        return None
    for lt in lesson_times:
        if lt.starttime == hhmm:
            try:
                return int(lt.number)
            except (TypeError, ValueError):
                return None
    return None


def detect_header(ws, max_scan: int = 8) -> Tuple[int, int, int, List[Tuple[int, str]]]:
    """Ищет строку шапки и колонки «Дни» / «Уроки» / классы."""
    for row in range(1, min(max_scan, ws.max_row) + 1):
        day_col = lesson_col = None
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row, col).value
            if v is None:
                continue
            key = norm(v)
            if key in {"дни", "день", "день недели", "дата", "weekday", "day"} and day_col is None:
                day_col = col
            elif key in {"уроки", "урок", "номер урока", "номер", "№", "период", "время",
                         "lesson", "period", "time"} and lesson_col is None:
                lesson_col = col
        if day_col is not None or lesson_col is not None:
            day_col = day_col or 1
            lesson_col = lesson_col or (day_col + 1)
            start = max(day_col, lesson_col) + 1
            classes = []
            for col in range(start, ws.max_column + 1):
                v = ws.cell(row, col).value
                if v is None or not str(v).strip():
                    continue
                classes.append((col, str(v).strip()))
            if classes:
                return row, day_col, lesson_col, classes
    raise ValueError("Не удалось найти шапку таблицы (ожидаются колонки «Дни», «Уроки» и названия классов).")


def read(path: str, sheet: Optional[str] = None, lesson_times=None) -> ExcelSchedule:
    """Читает workbook и возвращает список ячеек-занятий."""
    import openpyxl  # локальный импорт: библиотека нужна только для чтения Excel

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    sched = ExcelSchedule(path=path, sheet=ws.title, header_row=1, day_col=1, lesson_col=2)

    try:
        sched.header_row, sched.day_col, sched.lesson_col, sched.class_cols = detect_header(ws)
    except ValueError as exc:
        sched.warnings.append(str(exc))
        return sched

    merged_day = _merged_values(ws, sched.day_col)
    merged_lesson = _merged_values(ws, sched.lesson_col)

    current_day_raw = ""
    current_day = ""
    for row in range(sched.header_row + 1, ws.max_row + 1):
        day_val = merged_day.get(row, ws.cell(row, sched.day_col).value)
        if day_val is not None and str(day_val).strip():
            current_day_raw = str(day_val).strip()
            current_day = normalize_day(current_day_raw)
        lesson_val = merged_lesson.get(row, ws.cell(row, sched.lesson_col).value)
        period, time_text, period_raw = parse_period(lesson_val, lesson_times)
        if period is None and time_text is None:
            continue

        for col, class_name in sched.class_cols:
            value = ws.cell(row, col).value
            if value is None or not str(value).strip():
                continue
            pairs = parse_cell(value, ws.cell(row, col).coordinate, row, col)
            if not pairs:
                sched.warnings.append(f"{ws.title}!{ws.cell(row, col).coordinate}: пустое или "
                                      f"нераспознанное значение «{value}»")
                continue
            sched.entries.append(Entry(
                day_raw=current_day_raw, day=current_day, period_raw=period_raw,
                period=period, time_text=time_text, excel_class=class_name,
                col_index=col, row_index=row, pairs=pairs, raw=str(value).strip(),
            ))
    wb.close()

    if not sched.entries:
        sched.warnings.append("В таблице не найдено ни одной заполненной ячейки расписания.")
    unknown_days = sorted({e.day_raw for e in sched.entries if not e.day})
    if unknown_days:
        sched.warnings.append("Не распознаны названия дней: " + ", ".join(map(repr, unknown_days)))
    return sched


def _merged_values(ws, col: int) -> dict:
    """Значения объединённых ячеек колонки (день недели часто объединяют)."""
    out = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col <= col <= rng.max_col:
            value = ws.cell(rng.min_row, rng.min_col).value
            if value is not None:
                for r in range(rng.min_row, rng.max_row + 1):
                    out[r] = value
    return out


def normalize_day(text: str) -> str:
    """«Пн», «Понедельник», дата 14.09.2026 -> «понедельник»."""
    if text is None:
        return ""
    if isinstance(text, (_dt.date, _dt.datetime)):
        return WEEKDAY_NAMES[text.weekday()]
    key = norm(text)
    if key in DAY_WORDS:
        return DAY_WORDS[key]
    for word, name in DAY_WORDS.items():
        if key.startswith(word):
            return name
    m = re.search(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})", str(text))
    if m:
        try:
            d = _dt.date(int(m.group(3)) % 100 + 2000, int(m.group(2)), int(m.group(1)))
            return WEEKDAY_NAMES[d.weekday()]
        except ValueError:
            return ""
    return ""
