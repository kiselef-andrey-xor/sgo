# -*- coding: utf-8 -*-
"""Тесты ядра конвертера.

Запуск:  python tests/test_core.py            (без зависимостей)
     или:  python -m pytest tests -q
"""
from __future__ import annotations

import os
import sys
import tempfile
import xml.etree.ElementTree as ET

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE not in sys.path:
    sys.path.insert(0, BASE)

from converter import core, reference                                    # noqa: E402
from converter import excel_reader                                       # noqa: E402
from converter.aliases import Aliases                                    # noqa: E402
from converter.builder import build_timetable, week_name_from_date       # noqa: E402
from converter.excel_reader import normalize_day, parse_cell, parse_period  # noqa: E402
from converter.text_utils import fuzzy_score, norm, token_key            # noqa: E402
from converter import validate                                           # noqa: E402
from converter import reverse as reverse_mod                             # noqa: E402
from converter.teachers import TeacherIndex, parse_fio                   # noqa: E402

UPLOADS = os.path.join(BASE, "..", "uploads")
DEMO_XLSX = os.path.join(UPLOADS, "расписание по предметам (1).xlsx")
DEMO_XML = os.path.join(UPLOADS, "ExportCm.xml")
DEMO_XML2 = os.path.join(UPLOADS, "1_ExportCm.xml")
SUBJ_XLSX = os.path.join(UPLOADS, "расписание по предметам.xlsx")
TEACH_XLSX = os.path.join(UPLOADS, "расписание по учителям.xlsx")


# --------------------------------------------------------------------------- #
def test_text_utils():
    assert norm(" Химия_ ") == "химия"
    assert norm("спорт.зал2") == "спорт.зал2"
    assert token_key("Иностранный язык (английский)") == token_key("Иностранный (английский) язык")
    assert fuzzy_score("ДНКР", "ДНКР") == 1.0
    assert fuzzy_score("алгебра и начала анализа",
                       "Алгебра и начала математического анализа") > 0.8
    assert fuzzy_score("Практикум по химии", "Русский язык") < 0.4


def test_excel_cell_parsing():
    assert parse_cell("Русский язык (105)", "C2", 2, 3) == \
        parse_cell("Русский язык (105)", "C2", 2, 3)
    one = parse_cell("Русский язык (105)", "C2", 2, 3)[0]
    assert (one.subject, one.room) == ("Русский язык", "105")
    grp = parse_cell("Иностранный язык (английский) (305,307)", "O3", 3, 15)
    assert [(p.subject, p.room) for p in grp] == \
        [("Иностранный язык (английский)", "305"), ("Иностранный язык (английский)", "307")]
    two = parse_cell("Информатика, Труд (технология) (320,110)", "AR3", 3, 44)
    assert [(p.subject, p.room) for p in two] == [("Информатика", "320"), ("Труд (технология)", "110")]
    noroom = parse_cell("Классный час", "A10", 10, 3)
    assert [(p.subject, p.room) for p in noroom] == [("Классный час", "")]
    multi = parse_cell("Химия (317)\nФизика (318)", "AS45", 45, 45)
    assert [(p.subject, p.room) for p in multi] == [("Химия", "317"), ("Физика", "318")]
    # скобки предмета не принимаются за кабинет
    nop = parse_cell("Труд (технология)", "X1", 1, 2)
    assert [(p.subject, p.room) for p in nop] == [("Труд (технология)", "")]
    assert [(p.subject, p.room) for p in parse_cell("Физическая культура (спорт.зал)", "X2", 2, 2)] \
        == [("Физическая культура", "спорт.зал")]
    assert [(p.subject, p.room) for p in parse_cell("Физическая культура (БАС)", "X3", 3, 2)] \
        == [("Физическая культура", "БАС")]


def test_period_and_days():
    assert parse_period("3")[0] == 3
    assert parse_period(0)[0] == 0
    assert parse_period("1 урок")[0] == 1
    assert normalize_day("Пн") == "понедельник"
    assert normalize_day("Суббота") == "суббота"
    import datetime as dt
    assert normalize_day(dt.date(2026, 9, 14)) == "понедельник"


def test_reference_load():
    ref = reference.load(DEMO_XML)
    assert len(ref.classes) == 51
    assert len(ref.lesson_times) == 14
    assert ref.time_by_number(0).id == "1"
    assert [c.name for c in ref.classes_by_text("11и")] == ["11и (ест.-научн.)", "11и (техн.)"]
    assert ref.day_by_text("Сб").name == "суббота"
    assert ref.room_by_text("305").id == "25954"
    # 4б: все csg вложены без pclassid, но принадлежат классу
    c4b = ref.classes_by_text("4б")[0]
    assert c4b.csg and not c4b.csg[0].is_attached


def test_aliases_roundtrip():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "aliases.json")
        Aliases.default_file(path)
        al = Aliases.load(path)
        al.set("rooms", "Спорт. Зал", "Спортивный зал 1")
        al.set("subjects", "Физра", None)
        al.save()
        al2 = Aliases.load(path)
        assert al2.get("rooms", "спорт зал") == "Спортивный зал 1"
        assert Aliases.is_ignore(al2.get("subjects", "ФИЗРА"))
        al2.remove("rooms", "Спорт. Зал")
        assert al2.get("rooms", "спорт зал") is None


def test_week_name():
    import datetime as dt
    assert week_name_from_date(dt.date(2026, 9, 16)) == "Неделя 14.09.2026 - 20.09.2026"


def test_convert_demo():
    """Полный проход по реальным файлам из папки uploads."""
    if not (os.path.exists(DEMO_XLSX) and os.path.exists(DEMO_XML)):
        return
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "result.xml")
        res = core.convert(core.Options(
            excel_path=DEMO_XLSX, reference_path=DEMO_XML, output_path=out,
            aliases_path=os.path.join(BASE, "aliases.json"), write_output=True))
        assert res.ok
        assert res.stats["placements"] > 1000
        assert os.path.getsize(out) > 100_000

        root = ET.fromstring(res.xml_text)
        plan_ids = {c.get("id") for c in root.findall("Plan/class/csg")}
        room_ids = {r.get("id") for r in root.findall("Rooms/room")}
        csgs = root.findall("TimeTable/Week/Day/Lesson/csg")
        assert csgs and all(c.get("id") in plan_ids for c in csgs)
        assert all(not c.get("roomid") or c.get("roomid") in room_ids for c in csgs)
        # в одном уроке одна запись плана встречается один раз
        for day in root.findall("TimeTable/Week/Day"):
            for lesson in day.findall("Lesson"):
                ids = [c.get("id") for c in lesson.findall("csg")]
                assert len(ids) == len(set(ids))
        # словари из эталона не изменились
        orig = open(DEMO_XML, "rb").read().decode("windows-1251")
        assert orig[:orig.index("  <TimeTable>")] == res.xml_text[:res.xml_text.index("  <TimeTable>")]
        # кодировка windows-1251 и LF
        raw = open(out, "rb").read()
        assert raw.startswith(b'<?xml version="1.0" encoding="windows-1251"?>')
        assert b"\r\n" not in raw
        assert res.unresolved and res.mapping_csv and res.unresolved_csv


def test_unknown_room_noroom_mode():
    """Режим noroom: урок с неизвестным кабинетом остаётся, но без roomid."""
    if not (os.path.exists(DEMO_XLSX) and os.path.exists(DEMO_XML)):
        return
    with tempfile.TemporaryDirectory() as tmp:
        res = core.convert(core.Options(
            excel_path=DEMO_XLSX, reference_path=DEMO_XML,
            aliases_path=os.path.join(BASE, "aliases.json"),
            unknown_room_mode="noroom", write_output=False))
        no_room = [p for p in res.placements if p.room is None]
        assert no_room, "должны появиться уроки без кабинета"
        root = ET.fromstring(res.xml_text)
        plain = [c for c in root.findall("TimeTable/Week/Day/Lesson/csg")
                 if tuple(sorted(c.attrib)) == ("id",)]
        assert plain


def test_build_timetable_orders():
    ref = reference.load(DEMO_XML)
    res = core.convert(core.Options(excel_path=DEMO_XLSX, reference_path=DEMO_XML,
                                    aliases_path=os.path.join(BASE, "aliases.json"),
                                    write_output=False))
    block, notes = build_timetable(ref, res.placements)
    assert block.startswith("  <TimeTable>")
    assert block.rstrip().endswith("</TimeTable>")
    lines = block.split("\n")
    assert lines[1].startswith("    <Week ")
    # дни идут в порядке эталона, уроки — по возрастанию timeId
    days = [l for l in lines if l.strip().startswith("<Day ")]
    assert [d.split('name="')[1].split('"')[0] for d in days] == \
           [d.name for d in ref.days if any(p.day.id == d.id for p in res.placements)]


def test_validate_references():
    """Оба эталона целостны: битых ссылок и дублей id быть не должно."""
    for path in (DEMO_XML, DEMO_XML2):
        if not os.path.exists(path):
            continue
        findings = validate.validate(reference.load(path))
        errors = [f for f in findings if f.level == "error"]
        assert not errors, f"{path}: {[f.text for f in errors]}"


def test_teacher_fio_parsing():
    assert parse_fio("Кузнецова О.Вл.") == ("кузнецова", ["о", "вл"])
    assert parse_fio("Бочкарева Е..С.") == ("бочкарева", ["е", "с"])
    assert parse_fio("Лёшина Н.И.") == ("лешина", ["н", "и"])  # ё нормализуется в е
    if not os.path.exists(DEMO_XML2):
        return
    idx = TeacherIndex(reference.load(DEMO_XML2))
    # две Ольги Кузнецовы: «Вл.» — Владимировна, «В.» — неоднозначно
    tids = idx.resolve("Кузнецова О.Вл.")[0]
    assert len(tids) == 1 and idx.display(tids[0]) == "Кузнецова Ольга Владимировна"
    assert len(idx.resolve("Кузнецова О.В.")[0]) == 2
    assert len(idx.resolve("Бочкарева Е..С.")[0]) == 1
    assert idx.resolve("Киселев А.Ю.")[0] == []


def test_convert_with_teachers():
    """Две таблицы на входе: кабинеты сверяются, учителя уточняют выбор группы."""
    if not all(os.path.exists(p) for p in (SUBJ_XLSX, TEACH_XLSX, DEMO_XML2)):
        return
    res = core.convert(core.Options(
        excel_path=SUBJ_XLSX, teachers_path=TEACH_XLSX, reference_path=DEMO_XML2,
        aliases_path=os.path.join(BASE, "aliases.json"), write_output=False))
    assert res.stats["placements"] > 1400
    kinds = {i.kind for i in res.issues}
    # таблицы согласованы: расхождений кабинетов и числа групп быть не должно
    assert "grid_rooms_mismatch" not in kinds
    assert "grid_pairs_mismatch" not in kinds
    # протокол сопоставления содержит строки про учителей
    assert any(line.startswith("учитель;") for line in res.mapping_csv.splitlines())


def test_alias_dead_target_reported():
    """Правило с целью, которой нет в эталоне, не молчит, а попадает в отчёт;
    рабочее правило при этом применяется."""
    if not all(os.path.exists(p) for p in (SUBJ_XLSX, DEMO_XML2)):
        return
    import json as _json
    with tempfile.TemporaryDirectory() as tmp:
        al = os.path.join(tmp, "aliases.json")
        with open(al, "w", encoding="utf-8") as fh:
            _json.dump({
                "subjects": {}, "days": {}, "classes": {}, "teachers": {},
                "rooms": {"каб.пр.дея": "ЦПД",            # нет в эталоне
                          "112": "112",                  # нет в эталоне
                          "акт.зал": "Спортивный зал 1"},  # есть в эталоне
                "class_subjects": {},
            }, fh, ensure_ascii=False)
        res = core.convert(core.Options(
            excel_path=SUBJ_XLSX, reference_path=DEMO_XML2, aliases_path=al, write_output=False))
        dead = [i for i in res.issues if i.kind == "alias_target_missing"]
        assert {i.value for i in dead} >= {"каб.пр.дея -> ЦПД", "112 -> 112"}
        # рабочее правило применилось: уроки с «акт.зал» получили кабинет Спортивный зал 1
        sport = [p for p in res.placements if p.room_raw == "акт.зал"]
        assert sport and all(p.room is not None and p.room.name == "Спортивный зал 1" for p in sport)
        # «ЦПД» в результат не просочился
        assert all(p.room is None or p.room.name != "ЦПД" for p in res.placements)


def test_reverse_roundtrip():
    """XML -> XLSX -> разбор: сетка должна совпасть с расписанием из XML."""
    if not os.path.exists(DEMO_XML2):
        return
    import xml.etree.ElementTree as ET
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "rev.xlsx")
        res = reverse_mod.convert(reverse_mod.ReverseOptions(xml_path=DEMO_XML2, output_path=out))
        assert res.ok and os.path.getsize(out) > 20_000
        assert res.cells > 1000 and res.csg_total > 1000

        ref = reference.load(DEMO_XML2)
        root = ET.fromstring(ref.raw_text)
        number_by_time = {lt.id: lt.number for lt in ref.lesson_times}
        expect = []
        for day_el in root.findall("./TimeTable/Week/Day"):
            for lesson in day_el.findall("Lesson"):
                per = int(number_by_time.get(lesson.get("timeId"), "-1"))
                for c_el in lesson.findall("csg"):
                    csg = ref.plan.get(c_el.get("id"))
                    if not csg:
                        continue
                    cls = ref.classes.get(csg.owner_classid) or ref.classes.get(csg.pclassid)
                    room = ref.rooms.get(c_el.get("roomid"))
                    expect.append((day_el.get("name"), per, cls.name if cls else "?",
                                   csg.base_name, room.name if room else ""))
        sched = excel_reader.read(out)
        got = []
        for e in sched.entries:
            for p_ in e.pairs:
                got.append((e.day, e.period, e.excel_class, p_.subject, p_.room))
        assert sorted(got) == sorted(expect), (
            len(got), len(expect), sorted(set(got) ^ set(expect))[:5])
        # лист по учителям тоже читается и содержит ФИО с инициалами
        import openpyxl
        wb = openpyxl.load_workbook(out, read_only=True)
        assert "Расписание по учителям" in wb.sheetnames
        ws = wb["Расписание по учителям"]
        found = [c.value for row in ws.iter_rows(min_row=2, max_row=12) for c in row
                 if isinstance(c.value, str) and "." in c.value and "(" in c.value]
        assert found
        wb.close()


def test_auto_residual_informatics():
    """«Информатика» 5–6 классов автоподбирается к «Программирование и алгоритмика»
    только при подтверждении учителем; без таблицы по учителям — только подсказка."""
    if not all(os.path.exists(p) for p in (SUBJ_XLSX, TEACH_XLSX, DEMO_XML2)):
        return
    with_t = core.convert(core.Options(
        excel_path=SUBJ_XLSX, teachers_path=TEACH_XLSX, reference_path=DEMO_XML2,
        aliases_path=os.path.join(BASE, "aliases.json"), write_output=False))
    auto = {i.value for i in with_t.issues if i.kind == "subject_auto_residual"}
    assert "Информатика -> Программирование и алгоритмика" in auto
    assert not any(v.startswith("Практикум по физике ->") for v in auto)
    assert not any(u.value == "Информатика" for u in with_t.unresolved)
    without = core.convert(core.Options(
        excel_path=SUBJ_XLSX, reference_path=DEMO_XML2,
        aliases_path=os.path.join(BASE, "aliases.json"), write_output=False))
    assert not any(i.kind == "subject_auto_residual" for i in without.issues)
    sug = [u.suggestion for u in without.unresolved if u.value == "Информатика"]
    assert sug == ["Программирование и алгоритмика"]


def test_cli_help():
    from converter import cli
    parser = cli.build_parser()
    args = parser.parse_args(["a.xlsx", "-r", "b.xml", "--room-mode", "noroom", "--dry-run"])
    assert args.room_mode == "noroom" and args.dry_run


def main() -> int:
    failed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"  ok   {name}")
            except AssertionError as exc:
                failed += 1
                print(f"  FAIL {name}: {exc}")
            except Exception as exc:  # noqa: BLE001
                failed += 1
                print(f"  ERR  {name}: {type(exc).__name__}: {exc}")
    print("ВСЕ ТЕСТЫ ПРОЙДЕНЫ" if not failed else f"ПАДЕНИЙ: {failed}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
